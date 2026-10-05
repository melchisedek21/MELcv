"""Export du texte balisé (# titre, ## section, - puce, **gras**) en .docx, .pdf ou .txt."""

import re
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from . import config

BLEU = (0x1F, 0x3A, 0x5F)


def blocs(texte):
    """Découpe en blocs séparés par des lignes vides."""
    out, cur = [], []
    for ligne in texte.splitlines():
        if ligne.strip():
            cur.append(ligne.rstrip())
        elif cur:
            out.append(cur)
            cur = []
    if cur:
        out.append(cur)
    return out


def segments(texte):
    """[(texte, est_gras)] à partir du balisage **gras**."""
    return [(s, i % 2 == 1) for i, s in enumerate(re.split(r"\*\*", texte)) if s]


def lignes_logiques(bloc):
    """Retourne des (type, texte) : h1, h2, puce ou para (phrases longues rejointes)."""
    out, plain = [], []

    def vider():
        if not plain:
            return
        longues = len(plain) > 1 and all(len(x) > 70 for x in plain)
        out.append(("para", (" " if longues else "\n").join(plain)))
        plain.clear()

    for ligne in bloc:
        if ligne.startswith("## "):
            vider()
            out.append(("h2", ligne[3:].strip()))
        elif ligne.startswith("# "):
            vider()
            out.append(("h1", ligne[2:].strip()))
        elif re.match(r"^\s*[-•*]\s+", ligne) and not ligne.startswith("**"):
            vider()
            out.append(("puce", re.sub(r"^\s*[-•*]\s+", "", ligne)))
        else:
            plain.append(ligne.strip())
    vider()
    return out


def exporter_docx(texte, type_doc, chemin):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    cv = type_doc == "cv"
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.59), Cm(27.94)
    if cv:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.9)
    else:  # marges du canevas fourni
        s.top_margin, s.right_margin = Cm(1.2), Cm(2.34)
        s.bottom_margin, s.left_margin = Cm(2.0), Cm(2.0)
    taille = 10 if cv else 11
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Tahoma", Pt(taille)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    normal.paragraph_format.space_after = Pt(3 if cv else 6)

    def runs(p, t, gras=False, size=None, color=None):
        for seg, g in segments(t):
            r = p.add_run(seg)
            r.bold = g or gras
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor(*color)

    for bloc in blocs(texte):
        for kind, t in lignes_logiques(bloc):
            if kind == "h1":
                p = doc.add_paragraph()
                runs(p, t, True, taille + 8, BLEU)
                p.paragraph_format.space_after = Pt(2)
            elif kind == "h2":
                p = doc.add_paragraph()
                runs(p, t.upper(), True, taille + 1, BLEU)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                ppr = p._p.get_or_add_pPr()
                bordure = ppr.makeelement(qn("w:pBdr"), {})
                bas = bordure.makeelement(
                    qn("w:bottom"),
                    {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): "1F3A5F"},
                )
                bordure.append(bas)
                ppr.append(bordure)
            elif kind == "puce":
                p = doc.add_paragraph(style="List Bullet")
                runs(p, t)
                p.paragraph_format.space_after = Pt(1)
            else:
                p = doc.add_paragraph()
                for i, ligne in enumerate(t.split("\n")):
                    if i:
                        p.add_run().add_break()
                    runs(p, ligne)
                if not cv:
                    p.paragraph_format.space_after = Pt(8 if len(t) > 120 else 4)
    doc.save(chemin)


def _libreoffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def exporter_pdf(texte, type_doc, chemin):
    """PDF identique au Word si LibreOffice est installé, sinon mise en page simple (ReportLab)."""
    soffice = _libreoffice()
    if soffice:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "doc.docx"
            exporter_docx(texte, type_doc, src)
            try:
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                    check=True, capture_output=True, timeout=120,
                )  # fmt: skip
                shutil.copy(Path(tmp) / "doc.pdf", chemin)
                return
            except Exception:
                pass
    _pdf_reportlab(texte, type_doc, chemin)


def _pdf_reportlab(texte, type_doc, chemin):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

    def sur(t):
        t = t.replace("\u2019", "'").replace("\u2013", "-").replace("\u2014", "-")
        t = t.encode("cp1252", "replace").decode("cp1252")
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def md(t):
        corps = "".join(f"<b>{sur(s)}</b>" if g else sur(s) for s, g in segments(t))
        return corps.replace("\n", "<br/>")

    cv = type_doc == "cv"
    n = 9.5 if cv else 11
    normal = ParagraphStyle("n", fontName="Helvetica", fontSize=n, leading=n * 1.35, spaceAfter=4 if cv else 7)
    h1 = ParagraphStyle("h1", parent=normal, fontName="Helvetica-Bold", fontSize=n + 8, leading=n + 11)
    h2 = ParagraphStyle("h2", parent=normal, fontName="Helvetica-Bold", fontSize=n + 1, spaceBefore=8, spaceAfter=2)
    puce = ParagraphStyle("b", parent=normal, leftIndent=14, bulletIndent=3, spaceAfter=1.5)
    doc = SimpleDocTemplate(
        str(chemin), pagesize=LETTER, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=1.5 * cm, bottomMargin=1.8 * cm,
    )  # fmt: skip
    flow = []
    for bloc in blocs(texte):
        for kind, t in lignes_logiques(bloc):
            if kind == "h1":
                flow.append(Paragraph(md(t), h1))
            elif kind == "h2":
                flow += [Paragraph(md(t.upper()), h2), HRFlowable(width="100%", thickness=0.6, spaceAfter=3)]
            elif kind == "puce":
                flow.append(Paragraph(md(t), puce, bulletText="•"))
            else:
                flow.append(Paragraph(md(t), normal))
        if not cv:
            flow.append(Spacer(1, 2))
    doc.build(flow)


def nom_fichier(type_doc, nom_base):
    prefixe = {"cv": "CV", "motivation": "Lettre-motivation", "presentation": "Lettre-presentation"}[type_doc]
    base = re.sub(r"[^A-Za-z0-9À-ÿ]+", "-", nom_base).strip("-")[:40]
    return f"{prefixe}_{base + '_' if base else ''}{date.today().isoformat()}"


def exporter(texte, type_doc, fmt, nom_base=""):
    """Écrit le fichier dans le dossier sorties et retourne son chemin."""
    chemin = config.SORTIES / f"{nom_fichier(type_doc, nom_base)}.{fmt}"
    if fmt == "docx":
        exporter_docx(texte, type_doc, chemin)
    elif fmt == "pdf":
        exporter_pdf(texte, type_doc, chemin)
    elif fmt == "txt":
        chemin.write_text(texte, encoding="utf-8")
    else:
        raise ValueError(f"Format inconnu : {fmt}")
    return chemin
