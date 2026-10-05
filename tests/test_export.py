from docx import Document

from melcv import export

LETTRE = (
    "Ottawa, le 5 octobre 2026\n\nMadame, Monsieur,\n\n**Objet : Candidature au poste X**\n\n"
    "Premier paragraphe.\n\nVeuillez agréer, Madame, Monsieur, mes salutations distinguées."
)
CV = "# Marie Tremblay\nArchitecte\n\n## Expérience\n**Poste** — Société | 2020\n- Dirigé un projet"


def test_lignes_logiques_reconnait_le_balisage():
    types = [t for t, _ in export.lignes_logiques(["# Nom", "## Section", "- puce", "texte"])]
    assert types == ["h1", "h2", "puce", "para"]


def test_phrases_longues_sur_plusieurs_lignes_sont_rejointes():
    ligne = "a" * 80
    assert export.lignes_logiques([ligne, ligne]) == [("para", f"{ligne} {ligne}")]


def test_docx_contient_le_texte_et_le_gras(maison):
    chemin = export.exporter(LETTRE, "motivation", "docx", "EDC")
    doc = Document(str(chemin))
    textes = [p.text for p in doc.paragraphs]
    assert "Madame, Monsieur," in textes
    objet = next(p for p in doc.paragraphs if p.text.startswith("Objet"))
    assert objet.runs[0].bold


def test_cv_docx_a_des_puces_et_des_titres(maison):
    doc = Document(str(export.exporter(CV, "cv", "docx")))
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    assert any(p.text == "EXPÉRIENCE" for p in doc.paragraphs)


def test_pdf_et_txt_sont_crees(maison):
    pdf = export.exporter(LETTRE, "presentation", "pdf")
    assert pdf.read_bytes().startswith(b"%PDF")
    assert export.exporter(CV, "cv", "txt").read_text(encoding="utf-8") == CV


def test_pdf_reportlab_sans_libreoffice(maison, monkeypatch):
    monkeypatch.setattr(export, "_libreoffice", lambda: None)
    assert export.exporter(CV, "cv", "pdf").read_bytes().startswith(b"%PDF")


def test_format_inconnu(maison):
    import pytest

    with pytest.raises(ValueError):
        export.exporter(CV, "cv", "odt")
